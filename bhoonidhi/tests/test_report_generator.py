"""
Unit tests for Bhoonidhi PowerPoint Report Generator.
Tests 16:9 widescreen layout, slide structure, table generation,
image embedding, placeholder handling, and slide count capping.
"""

import tempfile
import unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation

from bhoonidhi.report_generator import SatelliteReportGenerator


class TestReportGenerator(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output_dir = Path(self.temp_dir.name)

        # Create a dummy 100x100 RGB JPEG image for testing
        self.test_img_path = self.output_dir / "sample_test_browse.jpg"
        img = Image.new("RGB", (100, 100), color=(15, 52, 96))
        img.save(str(self.test_img_path), "JPEG")

        # Mock STAC items with QC properties
        self.mock_items = [
            {
                "id": "SCENE_001",
                "collection": "EOS-04_SAR-FRS2_L2A",
                "properties": {
                    "Platform": "EOS-04",
                    "Instrument": "SAR",
                    "Mode": "FRS2",
                    "CycleNumber": 98,
                    "Imaging_Date": "29-OCT-2026",
                    "Dumping_Date": "29-OCT-2026",
                    "Dumping_Orbit": 25983,
                    "SceneCenterLat": 15.5,
                    "SceneCenterLon": 75.2,
                    "Online": "Y",
                },
                "preview_status": "downloaded",
                "preview_file_path": str(self.test_img_path),
            },
            {
                "id": "SCENE_002",
                "collection": "EOS-04_SAR-FRS2_L2A",
                "properties": {
                    "Platform": "EOS-04",
                    "Instrument": "SAR",
                    "Mode": "FRS2",
                    "CycleNumber": 98,
                    "Imaging_Date": "29-OCT-2026",
                    "Dumping_Date": "29-OCT-2026",
                    "Dumping_Orbit": 25983,
                    "SceneCenterLat": 16.0,
                    "SceneCenterLon": 75.8,
                    "Online": "Y",
                },
                "preview_status": "unavailable",
                "preview_file_path": None,  # No image available
            },
            {
                "id": "SCENE_003",
                "collection": "EOS-04_SAR-FRS2_L2A",
                "properties": {
                    "Platform": "EOS-04",
                    "Instrument": "SAR",
                    "Mode": "FRS2",
                    "CycleNumber": 98,
                    "Imaging_Date": "29-OCT-2026",
                    "Dumping_Date": "29-OCT-2026",
                    "Online": "Y",
                },
                "preview_status": "downloaded",
                "preview_file_path": str(self.test_img_path),
            },
        ]

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pptx_widescreen_and_structure(self):
        out_file = self.output_dir / "daily_report.pptx"
        generator = SatelliteReportGenerator(
            collection_id="EOS-04_SAR-FRS2_L2A",
            target_date="2026-10-29",
            manifest={"total_items_found": 3, "previews_downloaded": 2, "previews_unavailable": 1},
            max_slides=None,
        )

        res_path = generator.generate_report(self.mock_items, out_file)
        self.assertTrue(res_path.exists())

        # Inspect generated presentation
        prs = Presentation(str(res_path))

        # Check 16:9 Widescreen aspect ratio (13.333 x 7.5 inches)
        width_in = prs.slide_width / 914400
        height_in = prs.slide_height / 914400
        self.assertAlmostEqual(width_in, 13.333, places=2)
        self.assertAlmostEqual(height_in, 7.5, places=2)

        # Expected slides: 1 Cover + 1 Overview + 1 Completeness + 3 Product + 1 Summary = 7 slides
        self.assertEqual(len(prs.slides), 7)

        # Slide 1: Cover Slide
        cover_texts = " ".join([shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame])
        self.assertIn("DAILY SATELLITE IMAGERY REPORT", cover_texts)
        self.assertIn("2026-10-29", cover_texts)

        # Slide 2: Overview Table
        overview_tables = [s for s in prs.slides[1].shapes if s.has_table]
        self.assertEqual(len(overview_tables), 1)

        # Slide 3: Completeness Assurance Slide
        completeness_texts = " ".join([shape.text_frame.text for shape in prs.slides[2].shapes if shape.has_text_frame])
        self.assertIn("COMPLETENESS ASSURANCE & AUDIT", completeness_texts)
        self.assertIn("COMPLETENESS VERIFIED", completeness_texts)

        # Slide 4: Product Slide 1 (With real image embedded)
        prod1_pics = [s for s in prs.slides[3].shapes if s.shape_type == 13]
        self.assertEqual(len(prod1_pics), 1)
        prod1_tables = [s for s in prs.slides[3].shapes if s.has_table]
        self.assertEqual(len(prod1_tables), 1)

        # Slide 5: Product Slide 2 (Unavailable image placeholder)
        prod2_pics = [s for s in prs.slides[4].shapes if s.shape_type == 13]
        self.assertEqual(len(prod2_pics), 0)
        prod2_texts = " ".join([shape.text_frame.text for shape in prs.slides[4].shapes if shape.has_text_frame])
        self.assertIn("PREVIEW UNAVAILABLE", prod2_texts)

        # Slide 7: Summary Slide
        summary_texts = " ".join([shape.text_frame.text for shape in prs.slides[6].shapes if shape.has_text_frame])
        self.assertIn("QUALITY CONTROL & PIPELINE SUMMARY", summary_texts)

    def test_max_slides_capping(self):
        out_file = self.output_dir / "capped_report.pptx"
        # Set max_slides = 1
        generator = SatelliteReportGenerator(
            collection_id="EOS-04_SAR-FRS2_L2A",
            target_date="2026-10-29",
            max_slides=1,
        )

        res_path = generator.generate_report(self.mock_items, out_file)
        prs = Presentation(str(res_path))

        # Expected slides: 1 Cover + 1 Overview + 1 Completeness + 1 Product (capped) + 1 Summary = 5 slides
        self.assertEqual(len(prs.slides), 5)


if __name__ == "__main__":
    unittest.main()

